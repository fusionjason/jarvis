import json
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from collections import defaultdict

DATA_PATH = r"C:\Users\fusio\OneDrive\Documents\Assistant\Jarvis\data\assignments.json"
OUT_PATH = r"C:\Users\fusio\OneDrive\Desktop\Fall_2026_Assignment_Checklist.docx"

TAGS = {
    "Principles of Microeconomics": "Micro",
    "POL American Government": "Gov",
    "Political Science": "PoliSci",
}

data = json.load(open(DATA_PATH, encoding="utf-8"))

by_date = defaultdict(list)
for a in data:
    by_date[a["due_date"]].append(a)

doc = docx.Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("Fall 2026 Assignment Checklist")
r.bold = True
r.font.size = Pt(16)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.add_run("Updated through 08/30/2026 \u2013 [Micro] / [Gov] / [PoliSci] \u2013 * = high priority")

doc.add_paragraph()

for date in sorted(by_date.keys()):
    heading = doc.add_paragraph()
    r = heading.add_run(date)
    r.bold = True
    r.font.size = Pt(12)

    for a in sorted(by_date[date], key=lambda x: x["class"]):
        box = "\u2611" if a["status"] == "done" else "\u2610"
        tag = TAGS.get(a["class"], a["class"])
        star = " *" if a.get("priority") == "high" else ""
        line = f"{box} [{tag}] {a['title']}{star}"
        doc.add_paragraph(line, style=None)

doc.save(OUT_PATH)
print("saved:", OUT_PATH)
